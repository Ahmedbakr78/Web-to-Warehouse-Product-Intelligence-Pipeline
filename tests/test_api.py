"""API tests using the FastAPI test client (TestClient) against the seeded warehouse."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from app.api.main import app
from app.api.security import create_access_token
from app.core.db import read_session

pytestmark = pytest.mark.integration


@pytest.fixture(scope="module")
def client(database):
    """A TestClient bound to a bootstrapped + seeded database."""
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture(scope="module")
def admin_token():
    return create_access_token(1, role="admin", email="admin@example.com")


@pytest.fixture(scope="module")
def viewer_token():
    return create_access_token(3, role="viewer", email="viewer@example.com")


def auth(token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {token}"}


#: Module-level handle for the tests that walk a stateful 2FA enrolment.
ADMIN_TOKEN = create_access_token(1, role="admin", email="admin@example.com")


# --------------------------------------------------------------------------------------
# System
# --------------------------------------------------------------------------------------
def test_health(client):
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] in {"ok", "degraded"}
    assert payload["database"]["connected"] is True
    assert len(payload["sources"]) >= 3


def test_readiness_and_meta(client):
    from app.models import table_count

    assert client.get("/api/v1/health/ready").json()["ready"] is True
    meta = client.get("/api/v1/meta").json()
    # Derived from the ORM so adding a table does not silently break this test.
    assert meta["tables"] == table_count()
    assert meta["compliance"]["respect_robots_txt"] is True
    assert len(meta["features"]) >= 10


def test_openapi_schema_is_complete(client):
    schema = client.get("/openapi.json").json()
    assert len(schema["paths"]) > 50
    assert "bearerAuth" in json_dumps(schema) or "HTTPBearer" in json_dumps(schema)


def json_dumps(value) -> str:
    import json

    return json.dumps(value)


def test_unknown_route_returns_json_404(client):
    response = client.get("/api/v1/does-not-exist")
    assert response.status_code == 404
    assert "error" in response.json()


# --------------------------------------------------------------------------------------
# Authentication
# --------------------------------------------------------------------------------------
def test_login_with_seeded_account(client):
    response = client.post(
        "/api/v1/auth/login", json={"email": "admin@example.com", "password": "Admin@12345"}
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["access_token"] and payload["refresh_token"]
    assert payload["user"]["role"] == "admin"
    assert "manage_users" in payload["user"]["permissions"]


def test_login_rejects_bad_credentials(client):
    response = client.post(
        "/api/v1/auth/login", json={"email": "admin@example.com", "password": "wrong-password"}
    )
    assert response.status_code == 401
    assert response.json()["error"] == "authentication_failed"


def test_login_validates_payload(client):
    response = client.post("/api/v1/auth/login", json={"email": "not-an-email", "password": "x"})
    assert response.status_code == 422
    assert response.json()["details"]["errors"]


def test_register_creates_a_viewer_and_signs_in(client):
    import uuid

    email = f"newcomer-{uuid.uuid4().hex[:8]}@example.com"
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "full_name": "New Comer", "password": "Str0ng!Passw0rd"},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["access_token"] and payload["refresh_token"]
    assert payload["user"]["email"] == email
    assert payload["user"]["role"] == "viewer"
    assert "manage_users" not in payload["user"]["permissions"]

    # The new account signs in with the same credentials.
    again = client.post("/api/v1/auth/login", json={"email": email, "password": "Str0ng!Passw0rd"})
    assert again.status_code == 200


def test_register_rejects_duplicates_and_weak_passwords(client):
    import uuid

    email = f"dupe-{uuid.uuid4().hex[:8]}@example.com"
    body = {"email": email, "full_name": "Dupe User", "password": "Str0ng!Passw0rd"}
    assert client.post("/api/v1/auth/register", json=body).status_code == 201
    assert client.post("/api/v1/auth/register", json=body).status_code == 409

    weak = dict(body, email=f"weak-{uuid.uuid4().hex[:8]}@example.com", password="password123")
    denied = client.post("/api/v1/auth/register", json=weak)
    assert denied.status_code == 422

    # No role escalation: extra fields are ignored, not honoured.
    escalate = client.post(
        "/api/v1/auth/register",
        json={
            "email": f"escalate-{uuid.uuid4().hex[:8]}@example.com",
            "full_name": "Escalate",
            "password": "Str0ng!Passw0rd",
            "role": "admin",
        },
    )
    assert escalate.status_code == 201
    assert escalate.json()["user"]["role"] == "viewer"


def test_register_respects_the_kill_switch(client, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(settings, "registration_enabled", False)
    response = client.post(
        "/api/v1/auth/register",
        json={"email": "nobody@example.com", "full_name": "No Body", "password": "Str0ng!Passw0rd"},
    )
    assert response.status_code == 403


def test_me_requires_a_token(client):
    assert client.get("/api/v1/auth/me").status_code == 401


def test_me_returns_the_profile(client, admin_token):
    payload = client.get("/api/v1/auth/me", headers=auth(admin_token)).json()
    assert payload["email"]
    assert payload["rows_per_page"] >= 5


def test_refresh_rotates_tokens(client):
    login = client.post(
        "/api/v1/auth/login", json={"email": "viewer@example.com", "password": "Viewer@12345"}
    ).json()
    refreshed = client.post("/api/v1/auth/refresh", json={"refresh_token": login["refresh_token"]})
    assert refreshed.status_code == 200
    # JWTs are deterministic per second, so assert usability instead of inequality.
    token = refreshed.json()["access_token"]
    assert client.get("/api/v1/auth/me", headers=auth(token)).status_code == 200


# --------------------------------------------------------------------------------------
# Products
# --------------------------------------------------------------------------------------
def test_product_list_is_paginated(client, admin_token):
    response = client.get("/api/v1/products?page=1&page_size=5", headers=auth(admin_token))
    payload = response.json()
    assert len(payload["items"]) <= 5
    assert payload["total"] > 5
    assert payload["pages"] >= 2
    assert payload["has_next"] is True


def test_product_filters_narrow_the_result_set(client, admin_token):
    all_products = client.get("/api/v1/products?page_size=1", headers=auth(admin_token)).json()["total"]
    filtered = client.get("/api/v1/products?page_size=1&in_stock=true", headers=auth(admin_token)).json()[
        "total"
    ]
    assert 0 < filtered <= all_products


def test_product_search_is_case_insensitive(client, admin_token):
    payload = client.get("/api/v1/products?q=SAMSUNG&page_size=10", headers=auth(admin_token)).json()
    names = " ".join(item["canonical_name"] for item in payload["items"]).lower()
    assert "samsung" in names or payload["total"] == 0


def test_product_facets_and_categories(client, admin_token):
    facets = client.get("/api/v1/products/facets", headers=auth(admin_token)).json()
    assert facets["categories"] and facets["sources"]
    tree = client.get("/api/v1/products/categories", headers=auth(admin_token)).json()
    assert isinstance(tree, list) and tree


def test_product_detail_includes_history(client, admin_token):
    listing = client.get("/api/v1/products?page_size=1", headers=auth(admin_token)).json()
    product_id = listing["items"][0]["product_id"]
    detail = client.get(f"/api/v1/products/{product_id}", headers=auth(admin_token)).json()
    assert detail["product_id"] == product_id
    assert isinstance(detail["history"], list)


def test_unknown_product_returns_404(client, admin_token):
    response = client.get("/api/v1/products/99999999", headers=auth(admin_token))
    assert response.status_code == 404
    assert response.json()["error"] == "product_not_found"


def test_price_history_endpoint(client, admin_token):
    listing = client.get("/api/v1/products?page_size=1", headers=auth(admin_token)).json()
    product_id = listing["items"][0]["product_id"]
    payload = client.get(f"/api/v1/products/{product_id}/history", headers=auth(admin_token)).json()
    assert payload["product_id"] == product_id
    assert payload["points"] >= 1
    assert "first" in payload and "last" in payload


def test_csv_export(client, admin_token):
    response = client.get("/api/v1/analytics/export/products.csv", headers=auth(admin_token))
    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]
    assert "product_id" in response.text.splitlines()[0]


def test_xlsx_export(client, admin_token):
    response = client.get("/api/v1/export/products.xlsx?limit=5", headers=auth(admin_token))
    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    assert response.headers["Content-Disposition"].endswith('.xlsx"')
    assert int(response.headers["X-Row-Count"]) >= 0
    from openpyxl import load_workbook

    sheet = load_workbook(filename=io.BytesIO(response.content)).active
    assert sheet.max_row >= 1
    assert "product_id" in [cell.value for cell in sheet[1]]


# --------------------------------------------------------------------------------------
# Changes, analytics, pipeline, quality
# --------------------------------------------------------------------------------------
def test_price_change_feed(client, admin_token):
    payload = client.get("/api/v1/changes/price?page_size=5", headers=auth(admin_token)).json()
    assert "items" in payload
    if payload["items"]:
        row = payload["items"][0]
        assert row["direction"] in {"increase", "decrease"}


def test_change_summary(client, admin_token):
    payload = client.get("/api/v1/changes/summary?days=365", headers=auth(admin_token)).json()
    assert "total_events" in payload and "timeline" in payload


def test_top_movers_are_sorted(client, admin_token):
    rows = client.get("/api/v1/changes/top-movers?limit=5", headers=auth(admin_token)).json()
    magnitudes = [abs(row["change_pct"]) for row in rows if row.get("change_pct") is not None]
    assert magnitudes == sorted(magnitudes, reverse=True)


def test_analytics_kpi(client, admin_token):
    payload = client.get("/api/v1/analytics/kpi", headers=auth(admin_token)).json()
    assert payload["latest"]["products"] > 0
    assert payload["counts"]["fact_price_snapshot"] > 0


def test_pipeline_runs_and_detail(client, admin_token):
    runs = client.get("/api/v1/pipeline/runs?page_size=3", headers=auth(admin_token)).json()
    assert runs["items"]
    run_id = runs["items"][0]["run_id"]
    detail = client.get(f"/api/v1/pipeline/runs/{run_id}", headers=auth(admin_token)).json()
    assert detail["run_id"] == run_id
    assert "dq" in detail


def test_quality_latest_and_rules(client, admin_token):
    latest = client.get("/api/v1/quality/latest", headers=auth(admin_token)).json()
    assert latest["total"] >= 0
    rules = client.get("/api/v1/quality/rules", headers=auth(admin_token)).json()
    assert len(rules) == 12


def test_catalog_summary(client, admin_token):
    payload = client.get("/api/v1/catalog/summary", headers=auth(admin_token)).json()
    assert "totals" in payload


def test_sources_registry(client, admin_token):
    sources = client.get("/api/v1/sources", headers=auth(admin_token)).json()
    codes = {source["code"] for source in sources}
    assert "local_demo" in codes


# --------------------------------------------------------------------------------------
# Query lab (read-only guard)
# --------------------------------------------------------------------------------------
def test_query_executes_a_select(client, admin_token):
    response = client.post(
        "/api/v1/queries/execute", headers=auth(admin_token), json={"sql": "SELECT 1 AS ok"}
    )
    assert response.status_code == 200
    assert response.json()["rows"] == [[1]]


def test_query_rejects_writes(client, admin_token):
    for statement in (
        "DELETE FROM dim_product",
        "UPDATE dim_product SET price = 1",
        "DROP TABLE dim_product",
    ):
        response = client.post("/api/v1/queries/execute", headers=auth(admin_token), json={"sql": statement})
        assert response.status_code == 422, statement


def test_query_rejects_multiple_statements(client, admin_token):
    response = client.post(
        "/api/v1/queries/execute",
        headers=auth(admin_token),
        json={"sql": "SELECT 1; DELETE FROM dim_product"},
    )
    assert response.status_code == 422


def test_viewer_may_query_but_not_write(client, viewer_token):
    """`query` is granted to every role; writes are not."""
    response = client.post("/api/v1/queries/execute", headers=auth(viewer_token), json={"sql": "SELECT 1"})
    assert response.status_code == 200
    assert (
        client.post(
            "/api/v1/queries/execute", headers=auth(viewer_token), json={"sql": "DELETE FROM dim_product"}
        ).status_code
        == 422
    )


# --------------------------------------------------------------------------------------
# Account and administration
# --------------------------------------------------------------------------------------
def test_profile_update_round_trip(client, admin_token):
    headers = auth(admin_token)
    before = client.get("/api/v1/users/me", headers=headers).json()
    updated = client.patch("/api/v1/users/me", headers=headers, json={"rows_per_page": 50}).json()
    assert updated["rows_per_page"] == 50
    client.patch("/api/v1/users/me", headers=headers, json={"rows_per_page": before["rows_per_page"]})


def test_viewer_cannot_reach_admin_endpoints(client, viewer_token):
    for path in ("/api/v1/users", "/api/v1/users/stats", "/api/v1/audit"):
        response = client.get(path, headers=auth(viewer_token))
        assert response.status_code == 403, path
    # Public settings are readable by anyone, but writing them is admin-only.
    assert client.get("/api/v1/settings", headers=auth(viewer_token)).status_code == 200
    assert (
        client.put(
            "/api/v1/settings/ui.default_theme", headers=auth(viewer_token), json={"value": "dark"}
        ).status_code
        == 403
    )


def test_saved_view_lifecycle(client, admin_token):
    headers = auth(admin_token)
    created = client.post(
        "/api/v1/saved-views",
        headers=headers,
        json={"name": "pytest-view", "entity": "products", "filters": {"in_stock": True}},
    )
    assert created.status_code == 201
    view_id = created.json()["view_id"]
    assert (
        client.post(f"/api/v1/saved-views/{view_id}/favorite", headers=headers).json()["is_favorite"] is True
    )
    assert client.delete(f"/api/v1/saved-views/{view_id}", headers=headers).status_code == 200


def test_alert_lifecycle(client, admin_token):
    headers = auth(admin_token)
    created = client.post(
        "/api/v1/alerts",
        headers=headers,
        json={"name": "pytest-alert", "metric": "price_change_pct", "operator": "lt", "threshold": -12},
    )
    assert created.status_code == 201
    alert_id = created.json()["alert_id"]
    assert client.delete(f"/api/v1/alerts/{alert_id}", headers=headers).status_code == 200


def test_change_password_validation(client, admin_token):
    response = client.post(
        "/api/v1/auth/change-password",
        headers=auth(admin_token),
        json={"current_password": "Admin@12345", "new_password": "weak"},
    )
    assert response.status_code == 422


def test_api_key_creation_and_revocation(client, admin_token):
    headers = auth(admin_token)
    profile = client.get("/api/v1/users/me", headers=headers).json()
    created = client.post(
        f"/api/v1/users/{profile['user_id']}/api-keys", headers=headers, json={"name": "pytest-key"}
    )
    assert created.status_code == 201
    payload = created.json()
    assert payload["api_key"].startswith("pip_")
    assert (
        client.delete(
            f"/api/v1/users/{profile['user_id']}/api-keys/{payload['key_id']}", headers=headers
        ).status_code
        == 200
    )


def test_pipeline_trigger_requires_permission(client, viewer_token):
    response = client.post(
        "/api/v1/pipeline/run/sync", headers=auth(viewer_token), json={"limit_per_source": 1}
    )
    assert response.status_code == 403


def test_pipeline_trigger_runs_for_admin(client, admin_token):
    response = client.post(
        "/api/v1/pipeline/run/sync",
        headers=auth(admin_token),
        json={"sources": ["local_demo"], "limit_per_source": 3, "skip_dq": True},
    )
    assert response.status_code == 200
    payload = response.json()
    assert payload["run_id"]
    assert payload["status"] in {"success", "partial"}


def test_error_payloads_are_consistent(client, admin_token):
    response = client.get("/api/v1/products/99999999", headers=auth(admin_token))
    body = response.json()
    assert set(body) >= {"error", "message", "details"}


# --------------------------------------------------------------------------------------
# Appearance preferences: the profile must accept exactly what the dashboard offers
# --------------------------------------------------------------------------------------
APPEARANCE_PATCHES = [
    {"theme": "midnight"},
    {"theme": "high-contrast"},
    {"motion": "none"},
    {"font_scale": "xl"},
    {"direction": "rtl"},
    {"accent": "teal"},
]


@pytest.mark.parametrize("patch", APPEARANCE_PATCHES)
def test_appearance_preferences_round_trip(client, admin_token, patch):
    response = client.patch("/api/v1/users/me", json=patch, headers=auth(admin_token))
    assert response.status_code == 200, response.text
    for key, value in patch.items():
        assert response.json()[key] == value


def test_invalid_appearance_values_are_rejected(client, admin_token):
    """A bad theme or accent is a 422, never a silently stored typo."""
    bad_theme = client.patch("/api/v1/users/me", json={"theme": "neon"}, headers=auth(admin_token))
    assert bad_theme.status_code == 422

    bad_accent = client.patch("/api/v1/users/me", json={"accent": "chartreuse"}, headers=auth(admin_token))
    assert bad_accent.status_code == 422
    assert "accent must be one of" in bad_accent.text


def test_appearance_palettes_match_frontend(client):
    """The API's accent allow-list and the dashboard's presets must not drift apart."""
    import re
    from pathlib import Path

    from app.api.schemas import ACCENT_PRESETS

    theme_file = Path(__file__).resolve().parents[1] / "frontend" / "src" / "lib" / "theme.ts"
    if not theme_file.exists():  # frontend may be excluded from a backend-only checkout
        return
    block = re.search(
        r"export const ACCENTS: Record<string, Accent> = \{(.*?)\n\}", theme_file.read_text(), re.S
    )
    assert block, "could not find the ACCENTS map in frontend/src/lib/theme.ts"
    frontend_keys = set(re.findall(r"^\s{2}(\w+):\s*\{", block.group(1), re.M))
    assert frontend_keys == set(ACCENT_PRESETS), (
        f"frontend/backend accent drift: only frontend={frontend_keys - set(ACCENT_PRESETS)}, "
        f"only backend={set(ACCENT_PRESETS) - frontend_keys}"
    )


def test_profile_exposes_every_appearance_axis(client, admin_token):
    profile = client.get("/api/v1/users/me", headers=auth(admin_token)).json()
    for axis in ("theme", "accent", "density", "motion", "direction", "font_scale"):
        assert axis in profile, f"profile is missing the '{axis}' appearance axis"


def test_readiness_really_probes_the_views(client):
    """`views` must not be hard-coded to `pass`."""
    payload = client.get("/api/v1/health/ready").json()
    assert payload["checks"]["views"] == "pass"
    assert payload["views_present"] == payload["views_expected"] > 0
    assert payload["views_error"] is None


def test_change_pct_filters_are_signed_bounds(client, admin_token):
    """`min_change_pct=-5` keeps falls, `max_change_pct=5` keeps rises."""
    falls = client.get("/api/v1/products?min_change_pct=-5&page_size=200", headers=auth(admin_token)).json()
    assert all((item.get("price_change_pct") or 0) <= -5 for item in falls["items"])

    rises = client.get("/api/v1/products?max_change_pct=5&page_size=200", headers=auth(admin_token)).json()
    assert all((item.get("price_change_pct") or 0) >= 5 for item in rises["items"])


def test_user_pagination_total_respects_filters(client, admin_token):
    """The filtered count must match the number of rows, not the whole table."""
    everyone = client.get("/api/v1/users?page_size=200", headers=auth(admin_token)).json()
    viewers = client.get("/api/v1/users?role=viewer&page_size=200", headers=auth(admin_token)).json()
    assert viewers["total"] <= everyone["total"]
    assert viewers["total"] == len(viewers["items"])
    assert all(item["role"] == "viewer" for item in viewers["items"])


def test_run_http_log_is_scoped_to_the_run(client, admin_token):
    """`/runs/{run_id}/http` must not return the global audit log."""
    latest = client.get("/api/v1/pipeline/runs/latest", headers=auth(admin_token)).json()
    run_id = latest.get("run_id")
    if not run_id:
        return
    rows = client.get(f"/api/v1/pipeline/runs/{run_id}/http", headers=auth(admin_token)).json()
    assert all(row["run_id"] == run_id for row in rows)


def test_rebuild_aggregates_endpoint_exists(client, admin_token):
    response = client.post("/api/v1/pipeline/rebuild-aggregates", headers=auth(admin_token))
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "refreshed"


def test_change_password_requires_the_current_password(client, admin_token):
    # A strong candidate password with the wrong current password must be rejected
    # as an authentication failure, not a validation error.
    wrong = client.post(
        "/api/v1/users/me/password",
        json={"current_password": "definitely-wrong", "new_password": "BrandNew!Pass123"},
        headers=auth(admin_token),
    )
    assert wrong.status_code == 401, wrong.text

    # A weak candidate is a validation error, caught before any credential check.
    weak = client.post(
        "/api/v1/users/me/password",
        json={"current_password": "Admin@12345", "new_password": "short"},
        headers=auth(admin_token),
    )
    assert weak.status_code == 422


def test_pagination_order_rejects_injection():
    """`Pagination.order` must sanitise a hostile sort expression."""
    from app.api.deps import Pagination

    hostile = Pagination(page=1, page_size=25, sort_by="name; DROP TABLE app_user--", sort_dir="asc")
    assert "DROP" not in hostile.order.upper()
    assert hostile.order == "created_at ASC"

    safe = Pagination(page=1, page_size=25, sort_by="v.last_seen_at", sort_dir="asc")
    assert safe.order == "v.last_seen_at ASC"


def test_version_is_consistent():
    """One canonical VERSION file feeds the API, the package and pyproject."""
    import json
    import re
    from pathlib import Path

    import tomllib

    from app.core.config import canonical_version

    root = Path(__file__).resolve().parents[1]
    version = canonical_version()
    assert re.fullmatch(r"\d+\.\d+\.\d+", version)

    packaging = tomllib.loads((root / "pyproject.toml").read_text())["project"]["version"]
    assert packaging == version, "pyproject.toml has drifted from VERSION"

    frontend = json.loads((root / "frontend" / "package.json").read_text())["version"]
    assert frontend == version, "frontend/package.json has drifted from VERSION"


# --------------------------------------------------------------------------------------
# Two-factor authentication and session management
# --------------------------------------------------------------------------------------
def _enrol_two_factor(client, token: str) -> tuple[str, list[str]]:
    """Walk the whole enrolment flow and return (secret, recovery_codes)."""
    from app.services import twofactor

    setup = client.post("/api/v1/account/2fa/setup", headers=auth(token))
    assert setup.status_code == 200, setup.text
    secret = setup.json()["secret"]
    assert setup.json()["otpauth_uri"].startswith("otpauth://totp/")

    activated = client.post(
        "/api/v1/account/2fa/activate",
        json={"secret": secret, "code": twofactor.current_code(secret)},
        headers=auth(token),
    )
    assert activated.status_code == 200, activated.text
    return secret, activated.json()["recovery_codes"]


def test_two_factor_enrolment_flow(client, admin_token):
    from app.services import twofactor

    status = client.get("/api/v1/account/2fa/status", headers=auth(admin_token)).json()
    assert status["enabled"] is False

    # A wrong code must not enrol anything.
    setup = client.post("/api/v1/account/2fa/setup", headers=auth(admin_token)).json()
    wrong = client.post(
        "/api/v1/account/2fa/activate",
        json={"secret": setup["secret"], "code": "000000"},
        headers=auth(admin_token),
    )
    assert wrong.status_code == 401
    assert client.get("/api/v1/account/2fa/status", headers=auth(admin_token)).json()["enabled"] is False

    secret, codes = _enrol_two_factor(client, admin_token)
    assert len(codes) == 8

    after = client.get("/api/v1/account/2fa/status", headers=auth(admin_token)).json()
    assert after["enabled"] is True
    assert after["recovery_codes_remaining"] == 8
    assert after["enrolled_at"] is not None

    # The stored secret is not the plain one.
    with read_session() as session:
        from app.models.app_users import AppUser

        user = session.get(AppUser, 1)
        assert user.totp_secret is not None
        assert secret not in user.totp_secret

    # A live code is accepted by the verify endpoint.
    ok = client.post(
        "/api/v1/account/2fa/verify",
        json={"code": twofactor.current_code(secret)},
        headers=auth(admin_token),
    )
    assert ok.json()["valid"] is True

    # Tearing it down needs a valid code too.
    disabled = client.post(
        "/api/v1/account/2fa/disable",
        json={"code": twofactor.current_code(secret)},
        headers=auth(admin_token),
    )
    assert disabled.status_code == 200
    assert client.get("/api/v1/account/2fa/status", headers=auth(admin_token)).json()["enabled"] is False


def test_login_requires_the_second_factor_once_enrolled(client):
    from app.services import twofactor

    secret, codes = _enrol_two_factor(client, ADMIN_TOKEN)

    # Password alone is no longer enough.
    assert (
        client.post(
            "/api/v1/auth/login", json={"email": "admin@example.com", "password": "Admin@12345"}
        ).status_code
        == 401
    )

    # A wrong code is refused.
    assert (
        client.post(
            "/api/v1/auth/login",
            json={"email": "admin@example.com", "password": "Admin@12345", "totp_code": "111111"},
        ).status_code
        == 401
    )

    # A live code gets in.
    ok = client.post(
        "/api/v1/auth/login",
        json={
            "email": "admin@example.com",
            "password": "Admin@12345",
            "totp_code": twofactor.current_code(secret),
        },
    )
    assert ok.status_code == 200
    assert ok.json()["session_key"]

    # A recovery code also works, and is consumed.
    recovery = client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.com", "password": "Admin@12345", "recovery_code": codes[0]},
    )
    assert recovery.status_code == 200
    again = client.post(
        "/api/v1/auth/login",
        json={"email": "admin@example.com", "password": "Admin@12345", "recovery_code": codes[0]},
    )
    assert again.status_code == 401

    client.post(
        "/api/v1/account/2fa/disable",
        json={"code": twofactor.current_code(secret)},
        headers=auth(ADMIN_TOKEN),
    )


def test_two_factor_lockout_is_auditable(client):
    """Three bad codes lock the second factor; the counter must survive the 401."""
    from app.services import twofactor

    secret, _codes = _enrol_two_factor(client, ADMIN_TOKEN)

    for _ in range(3):
        client.post(
            "/api/v1/auth/login",
            json={"email": "admin@example.com", "password": "Admin@12345", "totp_code": "111111"},
        )

    status = client.get("/api/v1/account/2fa/status", headers=auth(ADMIN_TOKEN)).json()
    assert status["locked"] is True
    assert status["attempts_remaining"] == 0

    # Even a correct code is refused while locked...
    locked = client.post(
        "/api/v1/auth/login",
        json={
            "email": "admin@example.com",
            "password": "Admin@12345",
            "totp_code": twofactor.current_code(secret),
        },
    )
    assert locked.status_code == 401
    assert locked.json()["details"]["locked"] is True

    client.post(
        "/api/v1/account/2fa/disable",
        json={"code": twofactor.current_code(secret)},
        headers=auth(ADMIN_TOKEN),
    )


def test_sessions_are_listed_and_revoked(client, admin_token):
    login = client.post(
        "/api/v1/auth/login", json={"email": "viewer@example.com", "password": "Viewer@12345"}
    )
    assert login.status_code == 200
    session_key = login.json()["session_key"]
    assert session_key

    payload = client.get("/api/v1/account/sessions", headers=auth(login.json()["access_token"])).json()
    assert payload["active"] >= 1
    assert any(row["session_key"] == session_key for row in payload["sessions"])

    revoked = client.post(
        "/api/v1/account/sessions/revoke",
        json={"session_key": session_key},
        headers=auth(login.json()["access_token"]),
    )
    assert revoked.status_code == 200

    after = client.get("/api/v1/account/sessions", headers=auth(login.json()["access_token"])).json()
    revoked_row = next(row for row in after["sessions"] if row["session_key"] == session_key)
    assert revoked_row["is_active"] is False
    assert revoked_row["revoked_reason"] == "revoked by user"

    # A revoked session can no longer be refreshed.
    refresh = client.post(
        "/api/v1/auth/refresh",
        json={"refresh_token": login.json()["refresh_token"], "session_key": session_key},
    )
    assert refresh.status_code == 401


def test_revoke_others_keeps_the_current_session(client, admin_token):
    first = client.post(
        "/api/v1/auth/login", json={"email": "viewer@example.com", "password": "Viewer@12345"}
    ).json()
    second = client.post(
        "/api/v1/auth/login", json={"email": "viewer@example.com", "password": "Viewer@12345"}
    ).json()
    assert first["session_key"] != second["session_key"]

    result = client.post(
        f"/api/v1/account/sessions/revoke-others?current={second['session_key']}",
        headers=auth(second["access_token"]),
    )
    assert result.status_code == 200

    rows = client.get("/api/v1/account/sessions", headers=auth(second["access_token"])).json()["sessions"]
    kept = next(row for row in rows if row["session_key"] == second["session_key"])
    dropped = next(row for row in rows if row["session_key"] == first["session_key"])
    assert kept["is_active"] is True
    assert dropped["is_active"] is False


def test_session_refresh_tokens_are_stored_hashed():
    """A database read must not be enough to replay a session."""
    import sqlalchemy as sa

    with read_session() as session:
        rows = session.execute(sa.text("SELECT refresh_hash FROM app_session")).all()
    assert rows, "expected at least one recorded session"
    for (digest,) in rows:
        assert digest and len(digest) == 64
        assert not digest.startswith("eyJ"), "a JWT must never be stored in the clear"
